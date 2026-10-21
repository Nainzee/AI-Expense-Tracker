"""Generate the sample bank / UPI statements used by tests and for trying the import feature.

Run: .venv/bin/python sample_statements/generate_statements.py

Outputs (all names, accounts and UTRs are fake):
- hdfc_aug_oct_2026.csv            HDFC savings XX1234, 01-Aug..17-Oct-2026 (CSV export layout)
- hdfc_oct_2026.xlsx               same account, 10-Oct..20-Oct-2026 (overlaps the CSV on 10..17 Oct)
- icici_aug_oct_2026.pdf           ICICI savings XX5678, text PDF with multi-line narrations
- phonepe_oct_2026.pdf             PhonePe app statement (shares UTRs with the HDFC CSV)
- icici_aug_oct_2026_protected.pdf ICICI PDF encrypted with AES-128, password: ICICI5678
"""

from __future__ import annotations

import csv
import os
from datetime import date, datetime

HERE = os.path.dirname(os.path.abspath(__file__))
PDF_PASSWORD = "ICICI5678"
_FIXED = datetime(2026, 10, 21, 9, 0, 0)

# (date, narration, reference, debit, credit)
HDFC = [
    ("2026-08-01", "NEFT CR-CITI0000002-ACME TECHNOLOGIES PVT LTD-SALARY AUG 2026", "CITIN26080100001", 0, 85000.00),
    ("2026-08-03", "UPI-SWIGGY-swiggy@axisbank-UTIB0000123-629003100001-Payment", "0000629003100001", 412.00, 0),
    ("2026-08-05", "UPI-NETFLIX-netflixupi.payu@hdfcbank-HDFC0000499-629005100001-Autopay", "0000629005100001", 649.00, 0),
    ("2026-08-07", "UPI-JIO-jio@sbi-SBIN0000011-629007100001-Recharge", "0000629007100001", 299.00, 0),
    ("2026-08-09", "UPI-RAHUL SHARMA-9876543210@ybl-SBIN0001234-629009100001-dinner split", "0000629009100001", 1500.00, 0),
    ("2026-08-12", "ATW-416021XXXXXX1234-S1AW0123-BANGALORE", "", 2000.00, 0),
    ("2026-08-14", "UPI-BESCOM-bescom@hdfcbank-HDFC0000001-629014100001-Electricity", "0000629014100001", 1840.00, 0),
    ("2026-08-18", "UPI-UBER INDIA-uber@hdfcbank-HDFC0000001-629018100001-Ride", "0000629018100001", 286.40, 0),
    ("2026-08-22", "UPI-ZOMATO-zomato@hdfcbank-HDFC0000001-629022100001-Payment", "0000629022100001", 638.00, 0),
    ("2026-08-25", "CREDIT CARD PAYMENT HDFC CC XX5555 AUTOPAY", "CC2608250001", 12450.00, 0),
    ("2026-08-28", "NEFT DR-ICIC0000123-RAGHAV J-NETBANK-own account transfer", "N240826000123", 10000.00, 0),
    ("2026-09-01", "NEFT CR-CITI0000002-ACME TECHNOLOGIES PVT LTD-SALARY SEP 2026", "CITIN26090100001", 0, 85000.00),
    ("2026-09-02", "UPI-BLINKIT-blinkit@hdfcbank-HDFC0000001-629102100001-Groceries", "0000629102100001", 754.00, 0),
    ("2026-09-05", "UPI-NETFLIX-netflixupi.payu@hdfcbank-HDFC0000499-629105100001-Autopay", "0000629105100001", 649.00, 0),
    ("2026-09-11", "UPI-SWIGGY-swiggy@axisbank-UTIB0000123-629111100001-Payment", "0000629111100001", 389.00, 0),
    ("2026-09-15", "UPI-AMAZON PAY-amazonpay@apl-UTIB0000100-629115100001-REFUND", "0000629115100001", 0, 499.00),
    ("2026-09-16", "UPI-SRI LAKSHMI GENERAL STORES-srilakshmi@okaxis-UTIB0000200-629116100001-Payment",
     "0000629116100001", 320.00, 0),
    ("2026-09-20", "UPI-BESCOM-bescom@hdfcbank-HDFC0000001-629120100001-Electricity", "0000629120100001", 1610.00, 0),
    ("2026-09-25", "CREDIT CARD PAYMENT HDFC CC XX5555 AUTOPAY", "CC2609250001", 9875.00, 0),
    ("2026-10-01", "NEFT CR-CITI0000002-ACME TECHNOLOGIES PVT LTD-SALARY OCT 2026", "CITIN26100100001", 0, 85000.00),
    ("2026-10-05", "UPI-NETFLIX-netflixupi.payu@hdfcbank-HDFC0000499-629205100001-Autopay", "0000629205100001", 649.00, 0),
    ("2026-10-08", "UPI-SWIGGY-swiggy@axisbank-UTIB0000123-629208100001-Payment", "0000629208100001", 455.00, 0),
    ("2026-10-11", "POS 416021XXXXXX1234 NAMMA METRO BMRCL", "", 40.00, 0),
    ("2026-10-11", "POS 416021XXXXXX1234 NAMMA METRO BMRCL", "", 40.00, 0),
    ("2026-10-12", "UPI-RAHUL SHARMA-9876543210@ybl-SBIN0001234-629212100001-movie", "0000629212100001", 600.00, 0),
    ("2026-10-13", "UPI-SWIGGY-swiggy@axisbank-UTIB0000123-629213100001-Payment", "0000629213100001", 512.00, 0),
    ("2026-10-15", "POS 416021XXXXXX1234 DIGITAL TECH HUB", "POS2610150001", 2591.28, 0),
    ("2026-10-16", "POS 416021XXXXXX1234 ARTISAN BISTRO", "POS2610160001", 1239.00, 0),
    ("2026-10-17", "UPI-APOLLO PHARMACY-apollopharmacy@icici-ICIC0000001-629017100001-Payment", "0000629017100001",
     465.36, 0),
    # XLSX only (after the CSV period)
    ("2026-10-19", "POS 416021XXXXXX1234 SUPERMART RETAIL", "POS2610190001", 1349.25, 0),
    ("2026-10-20", "UPI-ZEPTO-zepto@ybl-YESB0000001-629220100001-Payment", "0000629220100001", 367.00, 0),
]
HDFC_OPENING = 42350.00

ICICI = [
    ("2026-08-04", ["UPI/AVENUE SUPERMARTS DMART/", "dmart@icici/629304100001/Groceries"], 1862.50, 0),
    ("2026-08-10", ["UPI/SPOTIFY/spotify@icici/", "629310100001/Subscription"], 119.00, 0),
    ("2026-08-20", ["BIL/ONL/LIC OF INDIA/", "LIC PREMIUM POLICY 4455"], 3250.00, 0),
    ("2026-08-28", ["NEFT-HDFC0000001-RAGHAV J-", "own account transfer"], 0, 10000.00),
    ("2026-09-03", ["UPI/AIRTEL/airtel@icici/", "629403100001/Postpaid"], 599.00, 0),
    ("2026-09-10", ["UPI/SPOTIFY/spotify@icici/", "629410100001/Subscription"], 119.00, 0),
    ("2026-09-18", ["CREDIT CARD PAYMENT ICICI", "CC XX7788 BBPS"], 6420.00, 0),
    ("2026-09-26", ["UPI/MYNTRA/myntra@icici/", "629426100001/REFUND"], 0, 1299.00),
    ("2026-10-02", ["UPI/AVENUE SUPERMARTS DMART/", "dmart@icici/629502100001/Groceries"], 2104.00, 0),
    ("2026-10-10", ["UPI/SPOTIFY/spotify@icici/", "629510100001/Subscription"], 119.00, 0),
    ("2026-10-14", ["UPI/AIRTEL/airtel@icici/", "629514100001/Postpaid"], 599.00, 0),
]
ICICI_OPENING = 25000.00

# (date, phrase, DEBIT/CREDIT, amount, UTR)
PHONEPE = [
    ("2026-10-08", "Paid to Swiggy", "DEBIT", 455.00, "629208100001"),
    ("2026-10-13", "Paid to Swiggy", "DEBIT", 512.00, "629213100001"),
    ("2026-10-14", "Received from Rahul Sharma", "CREDIT", 300.00, "629214100009"),
    ("2026-10-17", "Paid to Apollo Pharmacy", "DEBIT", 465.36, "629017100001"),
    ("2026-10-18", "Paid to BookMyShow", "DEBIT", 520.00, "629218100007"),
]


def _ddmmyy(d: str) -> str:
    return date.fromisoformat(d).strftime("%d/%m/%y")


def _money(v: float) -> str:
    return f"{v:,.2f}"


def _hdfc_rows():
    bal = HDFC_OPENING
    out = []
    for d, narr, ref, dr, cr in HDFC:
        bal = round(bal - dr + cr, 2)
        out.append((d, narr, ref, dr, cr, bal))
    return out


HDFC_HEADER = ["Date", "Narration", "Chq./Ref.No.", "Value Dt", "Withdrawal Amt.", "Deposit Amt.", "Closing Balance"]


def write_csv(path: str) -> None:
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["HDFC BANK Ltd.", "", "", "", "", "", ""])
        w.writerow(["Account No : 50100012341234", "", "", "", "", "", ""])
        w.writerow(["Statement From : 01/08/2026 To : 17/10/2026", "", "", "", "", "", ""])
        w.writerow(HDFC_HEADER)
        for d, narr, ref, dr, cr, bal in _hdfc_rows():
            if d > "2026-10-17":
                continue
            w.writerow([_ddmmyy(d), narr, ref, _ddmmyy(d), f"{dr:.2f}" if dr else "", f"{cr:.2f}" if cr else "",
                        f"{bal:.2f}"])


def write_xlsx(path: str) -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Statement"
    ws.append(["HDFC BANK Ltd."])
    ws.append(["Account No : 50100012341234"])
    ws.append(HDFC_HEADER)
    for d, narr, ref, dr, cr, bal in _hdfc_rows():
        if not "2026-10-10" <= d <= "2026-10-20":
            continue
        dt = datetime.fromisoformat(d)
        ws.append([dt, narr, ref, dt, dr or None, cr or None, bal])
        ws.cell(ws.max_row, 1).number_format = "DD/MM/YYYY"
        ws.cell(ws.max_row, 4).number_format = "DD/MM/YYYY"
    wb.properties.created = wb.properties.modified = _FIXED
    wb.save(path)


def _pdf():
    from fpdf import FPDF

    pdf = FPDF(format="A4")
    pdf.set_creation_date(_FIXED)
    pdf.set_auto_page_break(True, margin=15)
    pdf.add_page()
    return pdf


def _icici_pdf():
    pdf = _pdf()
    pdf.set_font("Helvetica", "B", 13)
    pdf.cell(0, 8, "ICICI Bank - Statement of Transactions", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", size=9)
    pdf.cell(0, 6, "Account Number: XXXXXXXX5678    Period: 01-08-2026 to 20-10-2026", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)
    pdf.set_font("Helvetica", "B", 9)
    widths = (24, 86, 26, 26, 28)
    for wd, h in zip(widths, ("Date", "Particulars", "Withdrawal", "Deposit", "Balance")):
        pdf.cell(wd, 7, h, border="B")
    pdf.ln(7)
    pdf.set_font("Helvetica", size=9)
    pdf.cell(widths[0], 6, "")
    pdf.cell(widths[1], 6, "Opening Balance")
    pdf.cell(widths[2] + widths[3], 6, "")
    pdf.cell(widths[4], 6, _money(ICICI_OPENING), align="R")
    pdf.ln(6)
    bal = ICICI_OPENING
    for d, lines, dr, cr in ICICI:
        bal = round(bal - dr + cr, 2)
        pdf.cell(widths[0], 5, date.fromisoformat(d).strftime("%d-%m-%Y"))
        pdf.cell(widths[1], 5, lines[0])
        pdf.cell(widths[2], 5, _money(dr) if dr else "", align="R")
        pdf.cell(widths[3], 5, _money(cr) if cr else "", align="R")
        pdf.cell(widths[4], 5, _money(bal), align="R")
        pdf.ln(5)
        for extra in lines[1:]:
            pdf.cell(widths[0], 5, "")
            pdf.cell(widths[1], 5, extra)
            pdf.ln(5)
        pdf.ln(1)
    pdf.ln(4)
    pdf.cell(0, 6, "This is a computer generated statement.")
    return pdf


def write_icici(path: str, protected_path: str) -> None:
    from fpdf.enums import EncryptionMethod

    _icici_pdf().output(path)
    pdf = _icici_pdf()
    pdf.set_encryption(owner_password="owner-" + PDF_PASSWORD, user_password=PDF_PASSWORD,
                       encryption_method=EncryptionMethod.AES_128)
    pdf.output(protected_path)


def write_phonepe(path: str) -> None:
    pdf = _pdf()
    pdf.set_font("Helvetica", "B", 13)
    pdf.cell(0, 8, "PhonePe Transaction Statement", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", size=9)
    pdf.cell(0, 6, "Oct 01, 2026 - Oct 20, 2026    Phone: XXXXXX3210", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(3)
    for i, (d, phrase, kind, amt, utr) in enumerate(PHONEPE):
        dt = date.fromisoformat(d)
        pdf.set_font("Helvetica", "B", 10)
        pdf.cell(0, 6, f"{dt.strftime('%b %d, %Y')} {phrase} {kind} Rs. {amt:,.2f}", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", size=8)
        pdf.cell(0, 5, f"{10 + i:02d}:15 AM Transaction ID T{dt.strftime('%y%m%d')}10{i:02d}001",
                 new_x="LMARGIN", new_y="NEXT")
        pdf.cell(0, 5, f"UTR No. {utr}", new_x="LMARGIN", new_y="NEXT")
        pdf.cell(0, 5, ("Paid by" if kind == "DEBIT" else "Credited to") + " XXXXXXXX1234", new_x="LMARGIN",
                 new_y="NEXT")
        pdf.ln(2)
    pdf.output(path)


def main() -> None:
    write_csv(os.path.join(HERE, "hdfc_aug_oct_2026.csv"))
    write_xlsx(os.path.join(HERE, "hdfc_oct_2026.xlsx"))
    write_icici(os.path.join(HERE, "icici_aug_oct_2026.pdf"), os.path.join(HERE, "icici_aug_oct_2026_protected.pdf"))
    write_phonepe(os.path.join(HERE, "phonepe_oct_2026.pdf"))
    print("Sample statements written to", HERE)


if __name__ == "__main__":
    main()
